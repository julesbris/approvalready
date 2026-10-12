"""Rendering report templates to HTML, PDF and Word.

A template version's body is a Jinja template rendered in a sandbox with autoescaping, from
a plain-data context (no ORM objects, no callables). The result is placed inside a frame
that this module owns, so every generated document carries the required report metadata
whatever the template says: date, project reference, project status, review status, the
assumptions it relies on, missing information, sources, limitations and the disclaimer.
``check_required`` verifies those sections are present before anything is stored.

PDF comes from WeasyPrint with every external fetch refused (no file://, no http://). Word
comes from converting the same HTML, so all three formats say the same thing. Templates
should only use: h1-h3, p, ul, ol, li, table, tr, th, td, strong, b, em, i, br.
"""

from __future__ import annotations

import io
import re
from functools import cache
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from jinja2 import StrictUndefined
from jinja2.sandbox import SandboxedEnvironment

REQUIRED_SECTIONS = (
    "metadata",
    "review-status",
    "assumptions",
    "missing-information",
    "sources",
    "limitations",
    "disclaimer",
)

BRAND_SCHEME = "ar-brand:"

DISCLAIMER = (
    "This report applies reviewed rules to the answers you gave. It is general information, "
    "not legal or planning advice, and it is not a council decision. Confidence levels show "
    "how far each result is supported by sources our team has verified. Check anything "
    "important with the authority or a qualified professional before acting on it."
)

BRAND_DIR = Path(__file__).parent / "brand"

# The ApprovalReady.au design system (packages/ui/src/tokens.css): navy and white lead, teal
# only marks what is done or approved, Outfit headings, Figtree text, hairline borders. The
# fonts are fetched from this module's brand folder by ``_BrandFetcher`` (nothing else is),
# so a browser opening the HTML download falls back to its own sans-serif.
_FONTS = (
    ("Outfit", 600, "Outfit-600.woff2"),
    ("Outfit", 700, "Outfit-700.woff2"),
    ("Figtree", 400, "Figtree-400.woff2"),
    ("Figtree", 600, "Figtree-600.woff2"),
    ("Figtree", 700, "Figtree-700.woff2"),
    ("JetBrains Mono", 400, "JetBrainsMono-400.woff2"),
)
_FONT_FACES = "".join(
    f'@font-face {{ font-family: "{family}"; font-weight: {weight}; '
    f'src: url("{BRAND_SCHEME}{name}") format("woff2"); }}\n'
    for family, weight, name in _FONTS
)

_STYLE = (
    _FONT_FACES
    + """
@page { size: A4; margin: 20mm 16mm 22mm;
  @top-right { content: string(doc-title); font: 600 8pt "Figtree", sans-serif;
    color: #4b5b74; }
  @bottom-left { content: "ApprovalReady.au"; font: 600 8pt "Outfit", sans-serif;
    color: #012455; }
  @bottom-right { content: "Page " counter(page) " of " counter(pages);
    font: 8pt "Figtree", sans-serif; color: #4b5b74; } }
@page :first { @top-right { content: none; } }
body { font-family: "Figtree", "DejaVu Sans", Arial, sans-serif; font-size: 10pt;
  color: #0a1a33; line-height: 1.5; max-width: 800px; margin: 0 auto; padding: 16px; }
@media print { body { padding: 0; } }
h1, h2, h3 { font-family: "Outfit", "Figtree", "DejaVu Sans", sans-serif; color: #012455;
  font-weight: 600; line-height: 1.2; break-after: avoid; }
h1 { font-size: 22pt; font-weight: 700; margin: 0; string-set: doc-title content(); }
.cover h1 { color: #ffffff; }
h2 { font-size: 14pt; margin: 22px 0 8px; padding-bottom: 5px;
  border-bottom: 1px solid #d7dfea; }
h3 { font-size: 11.5pt; margin: 14px 0 6px; }
p { margin: 0 0 8px; }
ul, ol { margin: 0 0 8px; padding-left: 18px; }
li { margin: 0 0 3px; }
li::marker { color: #01387f; }
strong, b { font-weight: 600; color: #012455; }
table { border-collapse: separate; border-spacing: 0; width: 100%; margin: 6px 0 12px;
  border: 1px solid #d7dfea; border-radius: 10px; box-sizing: border-box; }
tr { break-inside: avoid; }
th, td { padding: 6px 9px; text-align: left; vertical-align: top;
  border-bottom: 1px solid #d7dfea; }
tr:last-child > th, tr:last-child > td { border-bottom: none; }
th { background: #f2f5f9; color: #012455; font-weight: 600; }
.cover { background: #012455; color: #ffffff; border-radius: 16px; padding: 18px 22px 20px;
  margin: 0 0 14px; }
.cover-logo svg { display: block; height: 26px; width: auto; margin: 0 0 22px; }
.cover-project { margin: 6px 0 0; font-size: 11pt; color: #ffffff; }
.cover-project .ref { font-family: "JetBrains Mono", monospace; font-size: 9pt;
  color: #cfe0f5; }
table.details { font-size: 8.5pt; border: none; border-radius: 0; margin: 0 0 12px; }
table.details th, table.details td { background: none; padding: 2px 8px 2px 0;
  border-bottom: none; }
table.details th { color: #4b5b74; font-weight: 600; width: 26%; }
.note { border: 1px solid #d7dfea; background: #e7eef8; border-radius: 10px;
  padding: 9px 12px; }
section[data-section="sources"] li, section[data-section="limitations"] li { font-size: 9pt; }
footer[data-section="disclaimer"] { margin-top: 20px; padding: 2px 14px 6px;
  background: #f2f5f9; border-radius: 16px; break-inside: avoid; }
footer[data-section="disclaimer"] h2 { border-bottom: none; margin-top: 12px; }
"""
)

_FRAME = """<!doctype html>
<html lang="en-AU">
<head>
<meta charset="utf-8">
<title>{{ meta.title }} - {{ meta.project_reference }}</title>
<meta name="description" content="{{ meta.title }} for project {{ meta.project_reference }}">
<style>{{ style }}</style>
</head>
<body>
<header data-section="metadata">
<div class="cover">
<div class="cover-logo">{{ logo }}</div>
<h1>{{ meta.title }}</h1>
<p class="cover-project">{{ meta.project_title }}
<span class="ref">{{ meta.project_reference }}</span></p>
</div>
<table class="details">
<tr><th>Project</th><td>{{ meta.project_title }} ({{ meta.project_reference }})</td></tr>
<tr><th>Project status</th><td>{{ meta.project_status }}</td></tr>
<tr><th>Assessment date</th><td>{{ meta.assessed_on }}</td></tr>
<tr><th>Generated</th><td>{{ meta.generated_at }}</td></tr>
<tr><th>Document</th><td>{{ meta.document_id }}</td></tr>
<tr><th>Template</th><td>{{ meta.template_key }} version {{ meta.template_version }}</td></tr>
<tr><th>Rules engine</th><td>{{ meta.engine_version }}, facts {{ meta.facts_hash }}</td></tr>
</table>
</header>
<section data-section="review-status">
<p class="note"><strong>Review status: {{ meta.review_status }}.</strong>
{{ meta.review_status_detail }}</p>
{% if review.reviewer %}
<p>Reviewer: {{ review.reviewer }}
{%- if review.decided_on %}, {{ review.decided_on }}{% endif %}.</p>
{% endif %}
{% if review.notes %}<p>Reviewer's notes: {{ review.notes }}</p>{% endif %}
{% if review.changes %}
<h3>Changes made by the reviewer</h3>
<p>The reviewer changed these findings. The original result is kept for reference.</p>
<table>
<tr><th>Finding</th><th>Was</th><th>Now</th><th>Reason</th></tr>
{% for c in review.changes %}<tr><td>{{ c.finding }}</td><td>{{ c.before }}</td>
<td>{{ c.after }}</td>
<td>{{ c.reason }}{% if c.source %} (source: {{ c.source }}){% endif %}</td></tr>
{% endfor %}
</table>
{% endif %}
</section>
<main>
{{ body }}
</main>
<section data-section="assumptions">
<h2>Assumptions</h2>
<p>This report assumes the answers below are accurate and complete.</p>
{% if assumptions %}
<table>
<tr><th>Question</th><th>Your answer</th></tr>
{% for a in assumptions %}<tr><td>{{ a.label }}</td><td>{{ a.value }}</td></tr>{% endfor %}
</table>
{% else %}<p>No answers were recorded.</p>{% endif %}
</section>
<section data-section="missing-information">
<h2>Missing information</h2>
{% if missing %}
<p>These answers would make the results more certain:</p>
<ul>{% for m in missing %}<li>{{ m }}</li>{% endfor %}</ul>
{% else %}<p>None identified.</p>{% endif %}
</section>
<section data-section="sources">
<h2>Sources</h2>
{% if sources %}
<ul>{% for s in sources %}<li>{{ s.citation }}: {{ s.document_title }}, {{ s.organisation_name }}
({{ s.verification }}{% if not s.in_force %}, not in force on the assessment date{% endif %}).
{{ s.url }}</li>{% endfor %}</ul>
{% else %}<p>No sources were cited.</p>{% endif %}
</section>
<section data-section="limitations">
<h2>Limitations</h2>
{% if limitations %}
<ul>{% for l in limitations %}<li>{{ l }}</li>{% endfor %}</ul>
{% else %}<p>None recorded for the rules applied.</p>{% endif %}
</section>
<footer data-section="disclaimer">
<h2>Important</h2>
<p>{{ disclaimer }}</p>
</footer>
</body>
</html>
"""


class RenderError(ValueError):
    pass


def sandbox() -> SandboxedEnvironment:
    return SandboxedEnvironment(autoescape=True, undefined=StrictUndefined, trim_blocks=True)


def render_html(template_body: str, context: dict[str, Any]) -> str:
    """The complete HTML document: the template's body inside the required frame."""
    env = sandbox()
    try:
        body = env.from_string(template_body).render(**context)
        html = env.from_string(_FRAME).render(
            **context,
            body=_markup(body),
            style=_markup(_STYLE),
            logo=_markup(_logo_svg()),
            disclaimer=DISCLAIMER,
        )
    except Exception as exc:  # template bugs must never produce a half-rendered report
        raise RenderError(f"template could not be rendered: {exc}") from exc
    check_required(html)
    return html


def brand_style() -> Any:
    """The report stylesheet, for other documents built in the same style."""
    return _markup(_STYLE)


def brand_logo() -> Any:
    return _markup(_logo_svg())


@cache
def _logo_svg() -> str:
    """The reversed horizontal lockup, as supplied, for the navy cover band."""
    svg = (BRAND_DIR / "ar-lockup-horizontal-reversed.svg").read_text()
    return svg[svg.index("<svg") :]


@cache
def _brand_asset(name: str) -> bytes | None:
    if "/" in name or "\\" in name or name.startswith(".") or not name.endswith(".woff2"):
        return None
    path = BRAND_DIR / name
    return path.read_bytes() if path.is_file() else None


def _markup(value: str) -> Any:
    from markupsafe import Markup

    return Markup(value)  # noqa: S704 - body was rendered by the sandbox with autoescaping


def check_required(html: str) -> None:
    present = set(re.findall(r'data-section="([a-z-]+)"', html))
    missing = [s for s in REQUIRED_SECTIONS if s not in present]
    if missing:
        raise RenderError(f"report is missing required sections: {', '.join(missing)}")


def html_to_pdf(html: str) -> bytes:
    from weasyprint import HTML

    pdf: bytes = HTML(string=html, url_fetcher=_brand_fetcher()).write_pdf()
    return pdf


def _brand_fetcher() -> Any:
    from weasyprint import URLFetcher
    from weasyprint.urls import URLFetcherResponse

    class BrandFetcher(URLFetcher):  # type: ignore[misc]
        """Serves the brand fonts and nothing else. An empty protocol allowlist refuses
        every other URL: no local files, no network."""

        def fetch(self, url: str, headers: Any = None) -> Any:
            if url.startswith(BRAND_SCHEME):
                body = _brand_asset(url.removeprefix(BRAND_SCHEME))
                if body is not None:
                    return URLFetcherResponse(url, body, {"Content-Type": "font/woff2"})
            return super().fetch(url, headers)

    return BrandFetcher(allowed_protocols=())


# --- HTML to Word ----------------------------------------------------------------------

_SKIP = {"head", "style", "script", "title", "svg"}
_WS = re.compile(r"\s+")


class _DocxBuilder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        from docx import Document

        self.doc: Any = Document()
        _brand_docx(self.doc)
        self.para: Any = None
        self.skip = 0
        self.bold = 0
        self.italic = 0
        self.lists: list[str] = []
        self.rows: list[list[tuple[str, bool]]] | None = None
        self.cell: list[str] | None = None
        self.cell_header = False

    def _paragraph(self, style: str | None = None) -> Any:
        self.para = self.doc.add_paragraph(style=style)
        return self.para

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIP:
            self.skip += 1
        elif self.skip:
            return
        elif tag in {"h1", "h2", "h3"}:
            self.para = self.doc.add_heading(level=int(tag[1]) - 1)  # h1 is the Title style
        elif tag == "p" and self.cell is None:
            self._paragraph()
        elif tag in {"ul", "ol"}:
            self.lists.append(tag)
        elif tag == "li":
            ordered = bool(self.lists) and self.lists[-1] == "ol"
            self._paragraph("List Number" if ordered else "List Bullet")
        elif tag == "table":
            self.rows = []
        elif tag == "tr" and self.rows is not None:
            self.rows.append([])
        elif tag in {"th", "td"} and self.rows is not None:
            self.cell = []
            self.cell_header = tag == "th"
        elif tag in {"strong", "b"}:
            self.bold += 1
        elif tag in {"em", "i"}:
            self.italic += 1
        elif tag == "br":
            if self.cell is not None:
                self.cell.append("\n")
            elif self.para is not None:
                self.para.add_run().add_break()

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP:
            self.skip = max(self.skip - 1, 0)
        elif self.skip:
            return
        elif tag in {"h1", "h2", "h3", "p", "li"}:
            self.para = None
        elif tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
        elif tag in {"th", "td"} and self.rows is not None and self.cell is not None:
            if self.rows:
                text = _WS.sub(" ", "".join(self.cell)).strip()
                self.rows[-1].append((text, self.cell_header))
            self.cell = None
        elif tag == "table" and self.rows is not None:
            self._table(self.rows)
            self.rows = None
            self.para = None
        elif tag in {"strong", "b"}:
            self.bold = max(self.bold - 1, 0)
        elif tag in {"em", "i"}:
            self.italic = max(self.italic - 1, 0)

    def _table(self, rows: list[list[tuple[str, bool]]]) -> None:
        rows = [r for r in rows if r]
        if not rows:
            return
        width = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        _table_borders(table)
        for i, row in enumerate(rows):
            for j, (text, header) in enumerate(row):
                cell = table.cell(i, j)
                run = cell.paragraphs[0].add_run(text)
                run.bold = header
                if header:
                    run.font.color.rgb = _rgb(NAVY)
                    _shade(cell, SUNK)
        self.doc.add_paragraph()  # breathing room, as in the PDF

    def handle_data(self, data: str) -> None:
        if self.skip:
            return
        if self.cell is not None:
            self.cell.append(data)
            return
        text = _WS.sub(" ", data)
        if self.para is None:
            if not text.strip():
                return
            self._paragraph()
        if not self.para.runs:
            text = text.lstrip()
        if text:
            run = self.para.add_run(text)
            run.bold = self.bold > 0 or None
            run.italic = self.italic > 0 or None


NAVY = "012455"
INK = "0A1A33"
MUTED = "4B5B74"
BORDER = "D7DFEA"
SUNK = "F2F5F9"


def _rgb(hex_colour: str) -> Any:
    from docx.shared import RGBColor

    return RGBColor.from_string(hex_colour)


def _brand_docx(doc: Any) -> None:
    """Word styles after the design system: Figtree text, navy Outfit headings, the logo in
    the page header and the brand name with page numbers in the footer. Word falls back to
    its own fonts where Outfit and Figtree aren't installed."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Figtree"
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Outfit"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(NAVY)
        fonts = style.element.rPr.rFonts
        for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(theme), None)  # a theme font would win over Outfit
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)  # the default Title has a blue rule under it

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(BRAND_DIR / "ar-lockup-horizontal.png"), height=Mm(6))
    footer = section.footer.paragraphs[0]
    brand = footer.add_run("ApprovalReady.au    Page ")
    brand.font.size = Pt(8)
    brand.font.color.rgb = _rgb(MUTED)
    page = footer.add_run()
    page.font.size = Pt(8)
    page.font.color.rgb = _rgb(MUTED)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        page._r.append(node)


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _table_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)
        borders.append(node)
    table._tbl.tblPr.append(borders)


def html_to_docx(html: str, *, title: str) -> bytes:
    builder = _DocxBuilder()
    builder.doc.core_properties.title = title
    builder.feed(html)
    builder.close()
    out = io.BytesIO()
    builder.doc.save(out)
    return out.getvalue()
